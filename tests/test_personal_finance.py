from __future__ import annotations

from openpyxl import Workbook
from sqlalchemy import select

from financial_hub.models import AnnualFinanceSnapshot, MonthlyFinanceRecord
from financial_hub.services.personal_finance import (
    AnnualFinanceInput,
    MonthlyFinanceInput,
    calculate_yearly_summary,
    list_yearly_finance_rows,
    upsert_annual_snapshot,
    upsert_monthly_record,
)
from scripts.import_personal_finance import (
    import_personal_finance_workbook,
    read_personal_finance_workbook,
)


def test_monthly_and_annual_service_preserves_blanks_and_reports_coverage(db):
    _engine, factory, _ids = db
    with factory.begin() as session:
        upsert_monthly_record(session, MonthlyFinanceInput(2026, 1, 100, 40))
        upsert_monthly_record(session, MonthlyFinanceInput(2026, 2, 200, None))
        upsert_monthly_record(session, MonthlyFinanceInput(2026, 3, 0, 0))
        snapshot = upsert_annual_snapshot(
            session, AnnualFinanceInput(2026, 1_000, 2_000, 300)
        )
        summary = calculate_yearly_summary(session, 2026)
        rows = list_yearly_finance_rows(session)

    assert snapshot.total_asset_excluding_investment == 1_000
    assert summary.income == 300
    assert summary.expenditure == 40
    assert summary.income_months == 3
    assert summary.expenditure_months == 2
    assert summary.surplus is None
    assert rows[0].net_worth == 2_700

    with factory() as session:
        blank = session.scalar(
            select(MonthlyFinanceRecord).where(
                MonthlyFinanceRecord.year == 2026,
                MonthlyFinanceRecord.month == 2,
            )
        )
        assert blank is not None and blank.expenditure is None


def _write_reference_workbook(path):  # type: ignore[no-untyped-def]
    workbook = Workbook()
    monthly = workbook.active
    monthly.title = "Income & Expenditure"
    for year, start in ((2024, 5), (2025, 21), (2026, 37)):
        for index in range(12):
            row = start + index
            monthly.cell(row, 2).value = 100 + year - 2024
            monthly.cell(row, 6).value = -50
        # Leave the final three 2026 rows blank, as in the source workbook.
        if year == 2026:
            for row in range(46, 49):
                monthly.cell(row, 2).value = None
                monthly.cell(row, 6).value = None
            monthly.cell(45, 6).value = None
    annual = workbook.create_sheet("Balance Sheet")
    annual.cell(7, 2).value = "2024"
    annual.cell(8, 2).value = 95_527
    annual.cell(11, 2).value = 889_079
    annual.cell(15, 2).value = 50_000
    annual.cell(20, 2).value = 0
    workbook.save(path)


def test_workbook_import_is_fixed_layout_and_idempotent(db, tmp_path):
    _engine, factory, _ids = db
    source = tmp_path / "Personal Finance.xlsx"
    _write_reference_workbook(source)

    data = read_personal_finance_workbook(source)
    assert len(data.monthly) == 33
    assert len(data.annual) == 1
    assert data.annual[0].total_asset_excluding_investment == 145_527
    assert data.annual[0].total_portfolio_value == 889_079
    assert data.annual[0].total_debt == 0

    with factory.begin() as session:
        result = import_personal_finance_workbook(session, source)
    assert (result.imported_monthly, result.imported_annual) == (33, 1)

    with factory.begin() as session:
        result = import_personal_finance_workbook(session, source)
    assert (result.imported_monthly, result.imported_annual) == (0, 0)
    assert (result.skipped_monthly, result.skipped_annual) == (33, 1)

    with factory() as session:
        summary = calculate_yearly_summary(session, 2024)
        assert summary.income == 12 * 100
        assert summary.expenditure == 12 * 50
        assert session.scalar(select(AnnualFinanceSnapshot.total_debt)) == 0
