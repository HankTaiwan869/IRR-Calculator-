"""Import the simple Personal Finance fields from the reference workbook.

This is a one-time workbook-specific script.  It reads fixed cells from the
known workbook layout, skips blank future months, and never overwrites rows
already present in the database.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from financial_hub.database import (
    create_database_engine,
    default_database_path,
    session_factory,
)
from financial_hub.exceptions import ValidationError
from financial_hub.models import AnnualFinanceSnapshot, MonthlyFinanceRecord

DEFAULT_WORKBOOK = Path(
    r"C:\Users\hank8\OneDrive - NTHU\Others\Finance\Personal Finance.xlsx"
)

# Rows in the workbook's Income & Expenditure sheet containing Jan-Dec for
# each displayed fiscal year.  Income remains assigned to the displayed month.
MONTHLY_RANGES = ((2024, 5, 16), (2025, 21, 32), (2026, 37, 48))

# Balance Sheet values: header year row, then cash, portfolio, emergency funds,
# and debt rows.  The manual yearly asset field is cash + emergency funds.
ANNUAL_HEADER_ROW = 7
ANNUAL_CASH_ROW = 8
ANNUAL_PORTFOLIO_ROW = 11
ANNUAL_EMERGENCY_ROW = 15
ANNUAL_DEBT_ROW = 20


@dataclass(frozen=True, slots=True)
class WorkbookMonthlyRecord:
    year: int
    month: int
    income: int | None
    expenditure: int | None


@dataclass(frozen=True, slots=True)
class WorkbookAnnualSnapshot:
    year: int
    total_asset_excluding_investment: int | None
    total_portfolio_value: int | None
    total_debt: int | None


@dataclass(frozen=True, slots=True)
class WorkbookFinanceData:
    monthly: tuple[WorkbookMonthlyRecord, ...]
    annual: tuple[WorkbookAnnualSnapshot, ...]


@dataclass(frozen=True, slots=True)
class PersonalFinanceImportResult:
    imported_monthly: int
    skipped_monthly: int
    imported_annual: int
    skipped_annual: int


def _whole(value: object, field: str, *, absolute: bool = False) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValidationError(f"Workbook {field} must be a whole number.")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError) as error:
        raise ValidationError(f"Workbook {field} must be a whole number.") from error
    if not number.is_finite() or number != number.to_integral_value():
        raise ValidationError(f"Workbook {field} must be a whole number.")
    if absolute:
        number = abs(number)
    if number < 0:
        raise ValidationError(f"Workbook {field} cannot be negative.")
    return int(number)


def _sheet(workbook, name: str):  # type: ignore[no-untyped-def]
    try:
        return workbook[name]
    except KeyError as error:
        raise ValidationError(f"The workbook has no {name} sheet.") from error


def _read_monthly(sheet) -> list[WorkbookMonthlyRecord]:  # type: ignore[no-untyped-def]
    records: list[WorkbookMonthlyRecord] = []
    for year, start_row, end_row in MONTHLY_RANGES:
        rows = sheet.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=1,
            max_col=6,
            values_only=True,
        )
        for month, row in enumerate(rows, start=1):
            income = _whole(row[1], f"{year}-{month} income")
            # The workbook stores spending as a negative amount.
            expenditure = _whole(row[5], f"{year}-{month} expenditure", absolute=True)
            # Oct-Dec 2026 are blank layout rows and are not records.
            if income is None and expenditure is None:
                continue
            records.append(WorkbookMonthlyRecord(year, month, income, expenditure))
    return records


def _read_annual(sheet) -> list[WorkbookAnnualSnapshot]:  # type: ignore[no-untyped-def]
    records: list[WorkbookAnnualSnapshot] = []
    for column in range(2, sheet.max_column + 1):
        year_value = sheet.cell(ANNUAL_HEADER_ROW, column).value
        if year_value is None or str(year_value).strip() == "":
            continue
        try:
            year = int(str(year_value).strip())
        except (TypeError, ValueError) as error:
            raise ValidationError(
                "The Balance Sheet has an invalid year header."
            ) from error
        cash = _whole(sheet.cell(ANNUAL_CASH_ROW, column).value, f"{year} cash")
        portfolio = _whole(
            sheet.cell(ANNUAL_PORTFOLIO_ROW, column).value,
            f"{year} portfolio value",
        )
        emergency = _whole(
            sheet.cell(ANNUAL_EMERGENCY_ROW, column).value,
            f"{year} emergency funds",
        )
        debt = _whole(sheet.cell(ANNUAL_DEBT_ROW, column).value, f"{year} debt")
        if cash is None and portfolio is None and emergency is None and debt is None:
            continue
        assets = (
            None
            if cash is None and emergency is None
            else (cash or 0) + (emergency or 0)
        )
        records.append(WorkbookAnnualSnapshot(year, assets, portfolio, debt))
    return records


def read_personal_finance_workbook(source: Path | str) -> WorkbookFinanceData:
    """Read only the fixed monthly and annual fields from the source workbook."""

    path = Path(source).resolve()
    if not path.is_file():
        raise ValidationError("The Personal Finance workbook does not exist.")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        return WorkbookFinanceData(
            monthly=tuple(_read_monthly(_sheet(workbook, "Income & Expenditure"))),
            annual=tuple(_read_annual(_sheet(workbook, "Balance Sheet"))),
        )
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError(
            f"The Personal Finance workbook could not be read: {error}"
        ) from error
    finally:
        if workbook is not None:
            workbook.close()


def import_personal_finance_workbook(
    session,
    source: Path | str,  # type: ignore[no-untyped-def]
) -> PersonalFinanceImportResult:
    """Import workbook rows, skipping existing keys so reruns preserve edits."""

    workbook_data = read_personal_finance_workbook(source)
    imported_monthly = skipped_monthly = 0
    imported_annual = skipped_annual = 0
    for row in workbook_data.monthly:
        existing = session.scalar(
            select(MonthlyFinanceRecord).where(
                MonthlyFinanceRecord.year == row.year,
                MonthlyFinanceRecord.month == row.month,
            )
        )
        if existing is not None:
            skipped_monthly += 1
            continue
        session.add(
            MonthlyFinanceRecord(
                year=row.year,
                month=row.month,
                income=row.income,
                expenditure=row.expenditure,
            )
        )
        imported_monthly += 1
    for row in workbook_data.annual:
        existing = session.scalar(
            select(AnnualFinanceSnapshot).where(AnnualFinanceSnapshot.year == row.year)
        )
        if existing is not None:
            skipped_annual += 1
            continue
        session.add(
            AnnualFinanceSnapshot(
                year=row.year,
                total_asset_excluding_investment=row.total_asset_excluding_investment,
                total_portfolio_value=row.total_portfolio_value,
                total_debt=row.total_debt,
            )
        )
        imported_annual += 1
    session.flush()
    return PersonalFinanceImportResult(
        imported_monthly=imported_monthly,
        skipped_monthly=skipped_monthly,
        imported_annual=imported_annual,
        skipped_annual=skipped_annual,
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--database",
        type=Path,
        default=default_database_path(),
        help="SQLite database to import into (defaults to the app database).",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Personal Finance .xlsx workbook to read.",
    )
    args = parser.parse_args(argv)
    if not args.database.is_file():
        parser.error(f"Database does not exist: {args.database}")
    if not args.workbook.is_file():
        parser.error(f"Workbook does not exist: {args.workbook}")

    engine = create_database_engine(args.database)
    try:
        factory = session_factory(engine)
        with factory.begin() as session:
            result = import_personal_finance_workbook(session, args.workbook)
    except Exception as error:  # noqa: BLE001 - CLI boundary
        parser.error(str(error))
        return 2
    finally:
        engine.dispose()

    print(f"Imported monthly records: {result.imported_monthly}")
    print(f"Skipped monthly records: {result.skipped_monthly}")
    print(f"Imported annual records: {result.imported_annual}")
    print(f"Skipped annual records: {result.skipped_annual}")
    print(f"Database: {args.database.resolve()}")
    print(f"Workbook: {args.workbook.resolve()}")
    return 0


if __name__ == "__main__":  # pragma: no cover - exercised by the CLI
    raise SystemExit(main())
