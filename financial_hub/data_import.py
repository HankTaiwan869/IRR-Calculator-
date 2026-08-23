from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .exceptions import ValidationError
from .models import Portfolio, Security, TransactionKind
from .services.transactions import TransactionInput, create_transaction


@dataclass(frozen=True, slots=True)
class ImportResult:
    """Summary of a legacy file import."""

    imported_rows: int


@dataclass(frozen=True, slots=True)
class _ImportRow:
    row_id: object
    stock_code: object
    trade_date: object
    amount: object


def _round_integer(value: object) -> int:
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            raise InvalidOperation
        return int(number.quantize(Decimal(1), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError, TypeError) as error:
        raise ValidationError("Imported cash flow amount is invalid.") from error


def _read_sqlite_rows(path: Path) -> list[_ImportRow]:
    try:
        with sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True) as connection:
            table = connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='log'"
            ).fetchone()
            if table is None:
                raise ValidationError("The SQLite database has no log table.")
            return [
                _ImportRow(*row)
                for row in connection.execute(
                    "SELECT id, stock_code, time, amount FROM log ORDER BY id"
                ).fetchall()
            ]
    except ValidationError:
        raise
    except sqlite3.Error as error:
        raise ValidationError(f"The SQLite database could not be read: {error}") from error


def _read_excel_rows(path: Path) -> list[_ImportRow]:
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValidationError("The Excel workbook is empty.")
        positions = {
            str(header).strip().lower(): index
            for index, header in enumerate(headers)
            if header is not None
        }
        required = ("id", "stock_code", "date", "amount")
        missing = [name for name in required if name not in positions]
        if missing:
            raise ValidationError(
                "The Excel workbook is missing required columns: "
                + ", ".join(missing)
                + "."
            )
        return [
            _ImportRow(*(row[positions[name]] for name in required))
            for row in rows
            if any(cell is not None for cell in row)
        ]
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError(f"The Excel workbook could not be read: {error}") from error
    finally:
        if workbook is not None:
            workbook.close()


def _read_rows(path: Path) -> list[_ImportRow]:
    if path.suffix.lower() == ".xlsx":
        return _read_excel_rows(path)
    if path.suffix.lower() in {".db", ".sqlite", ".sqlite3"}:
        return _read_sqlite_rows(path)
    raise ValidationError("Choose an Excel (.xlsx) or SQLite (.db, .sqlite, .sqlite3) file.")


def _parse_date(value: object, row_id: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError as error:
        raise ValidationError(f"Imported row {row_id} has an invalid date.") from error


def import_transactions(session: Session, source: Path | str) -> ImportResult:
    """Import SQLite or Excel cash-flow data into the imported portfolio."""
    path = Path(source).resolve()
    if not path.is_file():
        raise ValidationError("Import file does not exist.")
    rows = _read_rows(path)

    portfolio = session.scalar(
        select(Portfolio).where(Portfolio.name == "Imported portfolio")
    )
    if portfolio is None:
        portfolio = Portfolio(name="Imported portfolio")
        session.add(portfolio)
        session.flush()

    for row in rows:
        symbol = str(row.stock_code).strip()
        if not symbol:
            raise ValidationError(f"Imported row {row.row_id} has no stock code.")
        security = session.scalar(select(Security).where(Security.symbol == symbol))
        if security is None:
            raise ValidationError(
                f"FinMind security master has no record for {symbol}; "
                "sync the security master before importing data."
            )
        create_transaction(
            session,
            TransactionInput(
                portfolio_id=portfolio.id,
                security_id=security.id,
                kind=TransactionKind.LEGACY_CASH_FLOW,
                trade_date=_parse_date(row.trade_date, row.row_id),
                amount=_round_integer(row.amount),
            ),
        )
    return ImportResult(imported_rows=len(rows))
